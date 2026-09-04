"""
Other Vertical SOPs Converter — DOCX (KMB/DIM Hawthorn SOPs) -> structured JSON
Replaces the old partially-converted JSON files with much richer, accurate records.

Table structure across all 4 docs:
  Table 0  : header info (client, org, version, address)
  Table 1  : version author table
  Table 2  : version history (Version Date | Version Control | Updated By | ... | Brief Description)
  Table 3  : HPS/Client contacts
  Table 4  : QWay contacts & escalation
  Table 5  : SLA/TAT
  Table 6  : Report communication OR Meetings
  Table 7  : Meetings OR Services checkboxes (varies per doc)
  Table 8  : Process/Specific Instructions (main SOP rules)
             Columns: SOW? | Specific | General Client Instruction | Received date | Received From
             OR: Process | Updates | Received date | Received from
  Table 9  : Holiday list
             Columns: Date | Voice Process | Non-Voice Process | Day of week
             OR: S.No | Observed Date | Holiday Name | Day | Status
"""

import os, re, json
from datetime import datetime
from pathlib import Path

TODAY = datetime.today().strftime('%Y-%m-%d')


def clean(s):
    return re.sub(r'[\s\u00a0]+', ' ', str(s or '')).strip()


def cell(row_cells, idx):
    return clean(row_cells[idx].text) if idx < len(row_cells) else ''


def make_record(practice_name, source_file, category, question, instruction,
                scope='General', status='ACTIVE', effective_date='NA',
                termination_date='NA', received_from='', billing_manager='',
                sop_reference=''):
    instruction = clean(instruction)
    if not instruction or len(instruction) < 5:
        return None
    return {
        'practice_name'   : practice_name,
        'scope'           : scope,
        'category'        : category,
        'question'        : question,
        'instruction'     : instruction,
        'effective_date'  : effective_date,
        'termination_date': termination_date,
        'status'          : status,
        'billing_manager' : billing_manager,
        'received_from'   : received_from,
        'sop_reference'   : sop_reference or f'{source_file} | {category}',
        'source_file'     : source_file,
        'last_updated'    : TODAY,
    }


def _detect_header(table):
    """Return list of lowered header strings from table row 0."""
    if not table.rows:
        return []
    return [clean(c.text).lower() for c in table.rows[0].cells]


def convert_docx(filepath):
    import docx as _docx
    doc = _docx.Document(filepath)
    source_file = os.path.basename(filepath)
    records = []

    # ── Derive practice name from doc header table (Table 0) ──────────────────
    practice_name = source_file
    org_name      = ''
    specialty     = ''
    version_str   = ''
    client_name   = ''

    if doc.tables:
        t0 = doc.tables[0]
        for row in t0.rows:
            for c in row.cells:
                txt = clean(c.text)
                if txt.lower().startswith('organization:'):
                    org_name = txt[13:].strip()
                elif txt.lower().startswith('client:'):
                    client_name = txt[7:].strip()
                elif txt.lower().startswith('specialty:'):
                    specialty = txt[10:].strip()
                elif txt.lower().startswith('version:'):
                    version_str = txt[8:].strip()

    if org_name:
        practice_name = org_name
    elif client_name:
        practice_name = client_name

    sop_ref_base = f"{source_file} | {version_str}" if version_str else source_file

    # ── Parse all paragraphs for section headings ──────────────────────────────
    # (used for contextual fallback)
    headings = []
    for p in doc.paragraphs:
        if p.text.strip() and ('heading' in p.style.name.lower() or p.text.strip().isupper()):
            headings.append(p.text.strip())

    # ── Parse each table ───────────────────────────────────────────────────────
    for ti, table in enumerate(doc.tables):
        if not table.rows:
            continue
        header = _detect_header(table)
        header_str = ' | '.join(header)

        # ── Version History table ───────────────────────────────────────────
        # Matches: "version date | version control | updated by | updated on | brief description..."
        # OR:      "version date | updated by | updated on | brief description..."
        if 'version date' in header_str and ('brief description' in header_str or 'updated by' in header_str):
            for row in table.rows[1:]:
                ver_date  = cell(row.cells, 0)
                # Find column indices dynamically
                upd_by_i  = next((i for i, h in enumerate(header) if 'updated by' in h), 2)
                desc_i    = next((i for i, h in enumerate(header) if 'brief description' in h), 4)
                approver_i= next((i for i, h in enumerate(header) if 'client approver' in h or 'client approval' in h), -1)

                upd_by   = cell(row.cells, upd_by_i)
                brief    = cell(row.cells, desc_i)
                approver = cell(row.cells, approver_i) if approver_i >= 0 else ''

                if not brief or brief.lower() in ('brief description of updates', ''):
                    continue

                r = make_record(
                    practice_name, source_file,
                    'Version History',
                    'What was recently updated in this SOP?',
                    f"[Version Date]: {ver_date}\n[Updated By]: {upd_by}\n"
                    f"[Brief Description]: {brief}"
                    + (f"\n[Client Approver]: {approver}" if approver else ''),
                    effective_date=ver_date or 'NA',
                    received_from=approver,
                    sop_reference=sop_ref_base + ' | Version History'
                )
                if r: records.append(r)
            continue

        # ── Contact Information (HPS/Client contacts) ─────────────────────
        if 'role' in header_str and ('name' in header_str or 'email' in header_str) and 'contact number' in header_str:
            # Determine if this is HPS or QWay contact table
            contact_type = 'HPS Contacts' if ti <= 3 else 'QWay Contacts'
            for row in table.rows[1:]:
                role   = cell(row.cells, 0)
                name   = cell(row.cells, 1)
                email  = cell(row.cells, 2) if len(row.cells) > 2 else ''
                phone  = cell(row.cells, 3) if len(row.cells) > 3 else ''
                if not name:
                    continue
                instr = (f"Role: {role}. Name: {name}."
                         + (f" Email: {email}." if email else '')
                         + (f" Phone: {phone}." if phone else ''))
                r = make_record(practice_name, source_file, contact_type,
                                f'Who is the {role} contact?', instr,
                                sop_reference=sop_ref_base + f' | {contact_type}')
                if r: records.append(r)
            continue

        # ── SLA / TAT table ───────────────────────────────────────────────
        if 'department' in header_str and 'tat' in header_str:
            for row in table.rows[1:]:
                dept = cell(row.cells, 0)
                tat  = cell(row.cells, 1) if len(row.cells) > 1 else ''
                if dept and tat:
                    r = make_record(practice_name, source_file, 'SLA & Turnaround Time',
                                    f'What is the TAT for {dept}?',
                                    f"Department: {dept}. Turnaround Time (TAT): {tat}.",
                                    sop_reference=sop_ref_base + ' | SLA')
                    if r: records.append(r)
            continue

        # ── Holiday List ──────────────────────────────────────────────────
        # Format 1: Date | Voice Process | Non-Voice Process | Day of week
        # Format 2: S.No | Observed Date | Holiday Name | Day | Status
        if ('voice process' in header_str or 'holiday name' in header_str
                or 'observed date' in header_str):
            holiday_records_voice = []
            holiday_records_nonvoice = []

            for row in table.rows[1:]:
                if 'voice process' in header_str:
                    date_val    = cell(row.cells, 0)
                    voice       = cell(row.cells, 1) if len(row.cells) > 1 else ''
                    non_voice   = cell(row.cells, 2) if len(row.cells) > 2 else ''
                    day_of_week = cell(row.cells, 3) if len(row.cells) > 3 else ''
                    if voice:
                        holiday_records_voice.append(f"{voice} ({date_val}, {day_of_week})")
                    if non_voice:
                        holiday_records_nonvoice.append(f"{non_voice} ({date_val}, {day_of_week})")
                elif 'holiday name' in header_str or 'observed date' in header_str:
                    date_val  = cell(row.cells, 1) if len(row.cells) > 1 else cell(row.cells, 0)
                    hol_name  = cell(row.cells, 2) if len(row.cells) > 2 else ''
                    day_val   = cell(row.cells, 3) if len(row.cells) > 3 else ''
                    status    = cell(row.cells, 4) if len(row.cells) > 4 else ''
                    if hol_name:
                        holiday_records_nonvoice.append(f"{hol_name} ({date_val}, {day_val})")

            if holiday_records_voice:
                r = make_record(practice_name, source_file, 'Holiday List',
                                'What are the US day shift (voice process) holidays?',
                                'Voice Process (US Day Shift) Holidays: ' + ', '.join(holiday_records_voice),
                                scope='Voice Process',
                                sop_reference=sop_ref_base + ' | Holiday List')
                if r: records.append(r)

            if holiday_records_nonvoice:
                r = make_record(practice_name, source_file, 'Holiday List',
                                'What are the non-voice process holidays?',
                                'Non-Voice Process (US Night Shift) Holidays: ' + ', '.join(holiday_records_nonvoice),
                                scope='Non-Voice Process',
                                sop_reference=sop_ref_base + ' | Holiday List')
                if r: records.append(r)

            # Combined all-holidays record
            all_holidays = list(dict.fromkeys(holiday_records_voice + holiday_records_nonvoice))
            if all_holidays:
                r = make_record(practice_name, source_file, 'Holiday List',
                                'What are all the holidays?',
                                'All Holidays 2026: ' + ', '.join(all_holidays),
                                sop_reference=sop_ref_base + ' | Holiday List')
                if r: records.append(r)
            continue

        # ── Main Process Specs / Client Instructions table ────────────────
        # Formats:
        #   [SOW | Specific | General Client Instruction | Received date | Received From]
        #   [Specific | General Client Instruction | Received date | Received From]
        #   [Process | Updates | Received date | Received from]
        has_sow      = 'sow' in header_str
        has_specific = 'specific' in header_str
        has_process  = 'process' in header_str
        has_updates  = 'updates' in header_str
        has_instr    = 'general client instruction' in header_str or 'updates' in header_str

        if (has_specific or has_process) and has_instr:
            sow_i   = next((i for i, h in enumerate(header) if h == 'sow'), -1)
            spec_i  = next((i for i, h in enumerate(header) if 'specific' in h or 'process' in h), 0)
            instr_i = next((i for i, h in enumerate(header) if 'instruction' in h or 'updates' in h), 1)
            date_i  = next((i for i, h in enumerate(header) if 'received date' in h or 'date' in h), 2)
            from_i  = next((i for i, h in enumerate(header) if 'received from' in h or 'from' in h), 3)

            for row in table.rows[1:]:
                sow       = cell(row.cells, sow_i) if sow_i >= 0 else ''
                specific  = cell(row.cells, spec_i)
                instr_txt = cell(row.cells, instr_i)
                recv_date = cell(row.cells, date_i)
                recv_from = cell(row.cells, from_i)

                if not instr_txt or instr_txt.lower() in ('general client instruction', 'updates', ''):
                    continue

                # Determine category from SOW column or specific column
                if sow:
                    cat = f"{sow} - Process Specifications"
                elif specific:
                    cat = 'Process Specifications'
                else:
                    cat = 'Process Specifications'

                # Build rich instruction text
                parts = []
                if specific:
                    parts.append(f"Topic: {specific}")
                parts.append(f"Instruction: {instr_txt}")

                r = make_record(
                    practice_name, source_file, cat,
                    specific or '',
                    '\n'.join(parts),
                    effective_date=recv_date or 'NA',
                    received_from=recv_from,
                    sop_reference=sop_ref_base + f' | {cat}'
                )
                if r: records.append(r)
            continue

        # ── Report Communication / Frequency ──────────────────────────────
        if 'report name' in header_str and 'frequency' in header_str:
            for row in table.rows[1:]:
                report_name = cell(row.cells, 0)
                description = cell(row.cells, 1) if len(row.cells) > 1 else ''
                frequency   = cell(row.cells, 2) if len(row.cells) > 2 else ''
                mode        = cell(row.cells, 3) if len(row.cells) > 3 else ''
                reviewer    = cell(row.cells, 4) if len(row.cells) > 4 else ''
                if not report_name:
                    continue
                instr = (f"Report: {report_name}. Description: {description}. "
                         f"Frequency: {frequency}. Mode: {mode}."
                         + (f" Reviewer: {reviewer}." if reviewer else ''))
                r = make_record(practice_name, source_file, 'Reports',
                                f'What is the {report_name} report?', instr,
                                sop_reference=sop_ref_base + ' | Reports')
                if r: records.append(r)
            continue

        # ── Report Communication (email/teams mode) ───────────────────────
        if 'item' in header_str and 'mode of communication' in header_str:
            for row in table.rows[1:]:
                item   = cell(row.cells, 0)
                mode   = cell(row.cells, 1) if len(row.cells) > 1 else ''
                origin = cell(row.cells, 2) if len(row.cells) > 2 else ''
                to_lst = cell(row.cells, 3) if len(row.cells) > 3 else ''
                freq   = cell(row.cells, 5) if len(row.cells) > 5 else ''
                if not item:
                    continue
                instr = (f"Communication Item: {item}. Mode: {mode}. "
                         f"Originator: {origin}. To: {to_lst}."
                         + (f" Frequency: {freq}." if freq else ''))
                r = make_record(practice_name, source_file, 'Report Communication',
                                f'How is {item} communicated?', instr,
                                sop_reference=sop_ref_base + ' | Communication')
                if r: records.append(r)
            continue

        # ── Meetings table ────────────────────────────────────────────────
        if 'item' in header_str and ('hps' in header_str or 'qway' in header_str or 'mode' in header_str):
            for row in table.rows[1:]:
                item   = cell(row.cells, 0)
                hps    = cell(row.cells, 1) if len(row.cells) > 1 else ''
                qway   = cell(row.cells, 2) if len(row.cells) > 2 else ''
                mode   = cell(row.cells, 3) if len(row.cells) > 3 else ''
                agenda = cell(row.cells, 4) if len(row.cells) > 4 else ''
                time   = cell(row.cells, 5) if len(row.cells) > 5 else ''
                if not item:
                    continue
                instr = (f"Meeting: {item}. HPS: {hps}. QWay: {qway}. "
                         f"Mode: {mode}. Agenda: {agenda}. Schedule: {time}.")
                r = make_record(practice_name, source_file, 'Meetings',
                                f'When is the {item}?', instr,
                                sop_reference=sop_ref_base + ' | Meetings')
                if r: records.append(r)
            continue

        # ── Services Scope checkboxes (Demo/Coding/Payments etc.) ─────────
        if 'demo' in header_str or 'charge entry' in header_str:
            # This is the services scope table — extract which services are in scope
            in_scope = []
            out_scope = []
            for row in table.rows[1:]:
                for i, h in enumerate(header):
                    v = cell(row.cells, i)
                    if v in ('✓', 'Yes', 'yes', 'TRUE', 'X'):
                        in_scope.append(h.title())
                    elif h and cell(row.cells, i) == '':
                        out_scope.append(h.title())
            if in_scope:
                r = make_record(practice_name, source_file, 'Services Scope',
                                'What services does QWay provide?',
                                f"Services in scope: {', '.join(in_scope)}."
                                + (f" Not in scope: {', '.join(out_scope)}." if out_scope else ''),
                                sop_reference=sop_ref_base + ' | Services')
                if r: records.append(r)
            continue

    return records


def convert_file(filepath, out_dir, department=None):
    fname = os.path.basename(filepath)
    records = convert_docx(filepath)

    if not records:
        print(f"  {fname}: 0 records")
        return 0

    # Deduplicate
    seen, deduped = set(), []
    for r in records:
        key = r['instruction'].strip().lower()[:200]
        if key not in seen:
            seen.add(key)
            deduped.append(r)

    # Extract practice name from first record
    practice_name = deduped[0]['practice_name'] if deduped else fname

    out_name = os.path.splitext(fname)[0] + '.json'
    out_path = os.path.join(out_dir, out_name)

    payload = {
        'source_file'  : fname,
        'practice_name': practice_name,
        'team'         : 'Other Vertical',
        'department'   : department or '',
        'converted_on' : TODAY,
        'records'      : deduped,
    }

    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)

    print(f"  {fname}: {len(deduped)} records -> {out_name}")
    return len(deduped)


def main():
    data_dir = os.path.join('data', 'Other Vertical SOPs')
    out_dir  = 'data_schema'
    os.makedirs(out_dir, exist_ok=True)

    files = list(Path(data_dir).glob('*.docx'))
    print(f"Converting {len(files)} Other Vertical SOP files...")
    total = 0
    for fp in sorted(files):
        total += convert_file(str(fp), out_dir)

    print(f"\nDone. Total records: {total}")


if __name__ == '__main__':
    main()
